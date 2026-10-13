"""Output files: Markdown, Word, PDF."""
import re
from datetime import datetime
from pathlib import Path

from .session import Chunk, Session, mask_secrets

FORMATS = ("md", "docx", "pdf")
_FONTS = Path(r"C:\Windows\Fonts")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")  # XML-illegal (tab/LF kept)


def _end(session: Session, ended: datetime | None) -> datetime:
    return ended or (session.latest().ts if session.chunks else datetime.now())


def _header(session: Session, ended: datetime) -> str:
    n = len(session.chunks)
    return (f"Gather — {session.started:%Y-%m-%d %H:%M}–{ended:%H:%M} · "
            f"{n} item{'' if n == 1 else 's'}")


def _clean(s: str) -> str:
    # lone surrogates (bad clipboard data) can't be encoded anywhere: swap for '?'
    return _ILLEGAL.sub("", s.encode("utf-8", "replace").decode())


def _label(i: int, c: Chunk, mask: bool = True) -> str:
    app, title = (mask_secrets(x) for x in (c.app, c.title)) if mask else (c.app, c.title)
    app = app[:-4] if app.lower().endswith(".exe") else app
    title = title if len(title) <= 60 else title[:59] + "…"
    who = " — ".join(p for p in (app, title) if p) or "unknown"
    return _clean(f"{i} · {who} · {c.ts:%H:%M}")


def _texts(session: Session, mask: bool) -> list[str]:
    return [(mask_secrets(c.text) if mask else c.text).encode("utf-8", "replace").decode()
            for c in session.chunks]


def render_markdown(session: Session, mask: bool = True, ended: datetime | None = None) -> str:
    """Whole session as Markdown text (also used for 'Copy last as text')."""
    body = "\n\n---\n\n".join(f"## {_label(i, c, mask)}\n\n{t}" for i, (c, t) in
                              enumerate(zip(session.chunks, _texts(session, mask)), 1))
    return f"# {_header(session, _end(session, ended))}\n\n{body}\n"


def _docx(session, path, mask, ended):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(_header(session, ended), 0)
    for i, (c, t) in enumerate(zip(session.chunks, _texts(session, mask)), 1):
        doc.add_heading(_label(i, c, mask), 2)
        for line in _ILLEGAL.sub("", t).split("\n"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    doc.save(str(path))


def _pdf(session, path, mask, ended):
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    ttf = all((_FONTS / f).exists() for f in ("segoeui.ttf", "segoeuib.ttf", "consola.ttf"))
    if ttf:
        pdf.add_font("Seg", "", str(_FONTS / "segoeui.ttf"))
        pdf.add_font("Seg", "B", str(_FONTS / "segoeuib.ttf"))
        pdf.add_font("Con", "", str(_FONTS / "consola.ttf"))
        fb = [f for f in ("seguisym.ttf", "msyh.ttc") if (_FONTS / f).exists()]
        for f in fb:
            pdf.add_font(f, "", str(_FONTS / f))
        pdf.set_fallback_fonts(fb)
        head, mono = "Seg", "Con"
    else:
        head, mono = "Helvetica", "Courier"

    def put(txt, font, style, size, h, fast=False):
        txt = _ILLEGAL.sub("", txt).replace("\t", "    ")
        if not ttf:
            txt = txt.encode("latin-1", "replace").decode("latin-1")
        pdf.set_font(font, style, size)
        # multi_cell line-breaking is ~15 s/MB; pure-ASCII monospace lines wrap by char count
        # and are drawn with text(). Non-ASCII keeps multi_cell (fallback fonts for CJK/emoji).
        cpl = max(1, int((pdf.w - pdf.l_margin - pdf.r_margin - 2 * pdf.c_margin) // pdf.get_string_width("x")))
        base = h / 2 + 0.3 * pdf.font_size  # baseline multi_cell uses inside a cell
        for line in txt.split("\n"):
            if fast and line.isascii() and line.isprintable():  # DEL/ctrl: font has no glyph
                for k in range(0, max(len(line), 1), cpl):
                    if pdf.y + h > pdf.page_break_trigger:
                        pdf.add_page()
                    if line:
                        pdf.text(pdf.l_margin + pdf.c_margin, pdf.y + base, line[k:k + cpl])  # c_margin: match multi_cell padding
                    pdf.set_xy(pdf.l_margin, pdf.y + h)
                continue
            # fpdf can't wrap an unbroken run when a line is wider than the page; wrap by char
            try:
                pdf.multi_cell(0, h, line or " ", new_x="LMARGIN", new_y="NEXT",
                               wrapmode="CHAR")
            except Exception:
                pdf.multi_cell(0, h, line.encode("ascii", "replace").decode() or " ",
                               new_x="LMARGIN", new_y="NEXT", wrapmode="CHAR")

    put(_header(session, ended), head, "B", 16, 9)
    for i, (c, t) in enumerate(zip(session.chunks, _texts(session, mask)), 1):
        pdf.ln(4)
        put(_label(i, c, mask), head, "B", 12, 7)
        put(t, mono, "", 9, 4.5, fast=True)
    pdf.output(str(path))


def write(session: Session, fmt: str, folder: Path, mask: bool = True,
          ended: datetime | None = None) -> Path:
    """Write gather-YYYY-MM-DD-HHMM.<fmt> into folder (created if missing;
    '-2', '-3' suffix on collision). Return the path."""
    if fmt not in FORMATS:
        raise ValueError(f"unknown format: {fmt}")
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    ended = ended or datetime.now()
    base = f"gather-{ended:%Y-%m-%d-%H%M}"
    path, n = folder / f"{base}.{fmt}", 2
    while path.exists():
        path, n = folder / f"{base}-{n}.{fmt}", n + 1
    if fmt == "md":
        path.write_bytes(render_markdown(session, mask, ended).encode("utf-8"))
    else:
        {"docx": _docx, "pdf": _pdf}[fmt](session, path, mask, ended)
    return path
